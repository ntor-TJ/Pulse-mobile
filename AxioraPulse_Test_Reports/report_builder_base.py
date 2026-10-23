import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Output Directories
OUTPUT_DIR = r"e:\Pulse\AxioraPulse_Test_Reports"
EVIDENCE_DIR = os.path.join(OUTPUT_DIR, "Evidence")
os.makedirs(OUTPUT_DIR, exist_ok=True)
os.makedirs(EVIDENCE_DIR, exist_ok=True)

# Styles
FONT_TITLE = Font(name="Calibri", size=16, bold=True, color="1F497D")
FONT_SUBTITLE = Font(name="Calibri", size=11, italic=True, color="595959")
FONT_SECTION = Font(name="Calibri", size=12, bold=True, color="1F497D")
FONT_HEADER = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
FONT_BOLD = Font(name="Calibri", size=11, bold=True)
FONT_REGULAR = Font(name="Calibri", size=11)
FONT_SMALL = Font(name="Calibri", size=10, italic=True, color="595959")

FILL_NAVY = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
FILL_STEEL = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
FILL_LIGHT_BLUE = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
FILL_ZEBRA = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

FILL_PASS = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
FONT_PASS = Font(name="Calibri", size=11, bold=True, color="375623")
FILL_FAIL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
FONT_FAIL = Font(name="Calibri", size=11, bold=True, color="C65911")
FILL_BLOCKED = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
FONT_BLOCKED = Font(name="Calibri", size=11, bold=True, color="833C0C")
FILL_PARTIAL = PatternFill(start_color="EDEDED", end_color="EDEDED", fill_type="solid")
FONT_PARTIAL = Font(name="Calibri", size=11, bold=True, color="595959")

BORDER_THIN = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)
BORDER_HEADER = Border(
    left=Side(style='thin', color='FFFFFF'),
    right=Side(style='thin', color='FFFFFF'),
    top=Side(style='medium', color='1F497D'),
    bottom=Side(style='medium', color='1F497D')
)

ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

def autofit(ws, max_len_cap=75):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or '')
            if cell.row in [1, 2]:
                continue
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), max_len_cap)

def style_header_row(ws, row_idx, num_cols):
    for c in range(1, num_cols + 1):
        cell = ws.cell(row=row_idx, column=c)
        cell.font = FONT_HEADER
        cell.fill = FILL_NAVY
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_HEADER

def style_data_row(ws, row_idx, num_cols, is_zebra=False, status_col=None):
    for c in range(1, num_cols + 1):
        cell = ws.cell(row=row_idx, column=c)
        cell.font = FONT_REGULAR
        cell.border = BORDER_THIN
        if is_zebra:
            cell.fill = FILL_ZEBRA
        if c == status_col and cell.value:
            v = str(cell.value).strip().upper()
            if "PASS" in v:
                cell.fill = FILL_PASS
                cell.font = FONT_PASS
            elif "FAIL" in v:
                cell.fill = FILL_FAIL
                cell.font = FONT_FAIL
            elif "BLOCK" in v:
                cell.fill = FILL_BLOCKED
                cell.font = FONT_BLOCKED
            elif "PARTIAL" in v:
                cell.fill = FILL_PARTIAL
                cell.font = FONT_PARTIAL

print("Base styling setup ready.")
